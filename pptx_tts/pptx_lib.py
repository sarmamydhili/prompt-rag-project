"""Shared PowerPoint narration, TTS, embed, and timing helpers."""

import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set
from zipfile import BadZipFile, ZipFile

from dotenv import load_dotenv
from lxml import etree
from mutagen.mp3 import MP3
from openai import OpenAI
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

BASE_DIR = Path(__file__).resolve().parent

TTS_MODEL = "gpt-4o-mini-tts"
NARRATION_MODEL = "gpt-4o-mini"
VOICE = "onyx"
AUDIO_FORMAT = "mp3"
SPEECH_INSTRUCTIONS = (
    "Speak clearly and naturally in a friendly instructional tone. "
    "Use a comfortable teaching pace. Pause naturally between ideas and "
    "slow down slightly for important concepts. Avoid sounding like you "
    "are simply reading text."
)
NARRATION_SYSTEM_PROMPT = (
    "You write spoken presenter scripts for AP Cybersecurity slide decks. "
    "Your audience is middle school and high school students who are new to "
    "cybersecurity. Be elaborate and detailed: explain terms, give context, "
    "and use simple examples. Stay faithful to the slide content, but you may "
    "add relevant background that helps understanding. Write as a natural "
    "spoken script — no bullet lists, no meta phrases like 'this slide shows'. "
    "Aim for about 150–250 words unless the slide is very dense. "
    "Do not exceed 3500 characters."
)
NARRATION_USER_TEMPLATE = (
    "Deck: {deck_name}\n"
    "Slide: {slide_num}\n\n"
    "Slide content:\n{content}\n"
    "{existing_notes_section}"
    "\nWrite the full presenter narration script for this slide."
)
MAX_TTS_INPUT_CHARS = 4096
MAX_NARRATION_CHARS = 3500

AUTO_PLAY = True
HIDE_ICON_DURING_SHOW = True
ICON_SIZE = Inches(0.5)
ICON_MARGIN = Inches(0.62)

BUFFER_MS = 500
DEFAULT_SLIDE_DURATION_MS = 3000
ALLOW_CLICK_ADVANCE = True

PLACEHOLDER_PATTERNS = [
    re.compile(r"^click to add notes\.?$", re.IGNORECASE),
    re.compile(r"^click to edit master text styles\.?$", re.IGNORECASE),
    re.compile(r"^click to edit notes\.?$", re.IGNORECASE),
]
NARRATION_LABEL_LINE = re.compile(r"^NARRATION:?\s*$", re.IGNORECASE)
NARRATION_LABEL_PREFIX = re.compile(r"^NARRATION:\s*", re.IGNORECASE)

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NSMAP = {"p": P_NS}
NSMAP_A = {"p": P_NS, "a": A_NS}

_media_sha1_patch_applied = False


def _patch_media_sha1_lookup() -> None:
    """Work around python-pptx crash when existing media parts lack sha1."""
    global _media_sha1_patch_applied
    if _media_sha1_patch_applied:
        return

    from pptx.package import _MediaParts

    def _find_by_sha1(self, sha1):
        for media_part in self:
            if not hasattr(media_part, "sha1"):
                continue
            if media_part.sha1 == sha1:
                return media_part
        return None

    _MediaParts._find_by_sha1 = _find_by_sha1
    _media_sha1_patch_applied = True


def safe_save_presentation(presentation: Presentation, pptx_path: Path) -> None:
    """Save pptx to a temp file and verify zip integrity before replacing."""
    pptx_path = pptx_path.resolve()
    tmp_path = pptx_path.with_name(f"{pptx_path.stem}_tmp.pptx")
    presentation.save(str(tmp_path))
    try:
        with ZipFile(tmp_path) as zf:
            bad_member = zf.testzip()
            if bad_member:
                raise BadZipFile(f"Corrupt member: {bad_member}")
    except BadZipFile as exc:
        tmp_path.unlink(missing_ok=True)
        raise RuntimeError(f"Save failed integrity check: {exc}") from exc
    shutil.move(str(tmp_path), str(pptx_path))


def _shape_has_audio(element) -> bool:
    if etree.ElementBase.xpath(element, ".//a:audioFile", namespaces=NSMAP_A):
        return True
    if etree.ElementBase.xpath(element, ".//p:videoFile", namespaces=NSMAP):
        return True
    return False


def remove_slide_media(slide) -> int:
    """Remove existing audio/media shapes and timing from a slide."""
    removed = 0
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.MEDIA or _shape_has_audio(shape._element):
            shape._element.getparent().remove(shape._element)
            removed += 1
        elif shape.name and shape.name.lower().endswith(".mp3"):
            shape._element.getparent().remove(shape._element)
            removed += 1

    slide_element = slide._element
    for timing in xpath(slide_element, ".//p:timing"):
        timing.getparent().remove(timing)

    return removed


def xpath(element, query: str):
    return etree.ElementBase.xpath(element, query, namespaces=NSMAP)


def strip_narration_label(text: str) -> str:
    if not text:
        return ""

    lines = text.splitlines()
    if not lines:
        return text

    if NARRATION_LABEL_LINE.match(lines[0].strip()):
        lines = lines[1:]
    elif NARRATION_LABEL_PREFIX.match(lines[0]):
        lines[0] = NARRATION_LABEL_PREFIX.sub("", lines[0], count=1)

    return "\n".join(lines)


def clean_narration(text: str) -> str:
    if not text:
        return ""

    text = strip_narration_label(text)
    lines = [line.strip() for line in text.splitlines()]
    cleaned_lines = []
    prev_blank = False

    for line in lines:
        if not line:
            if not prev_blank:
                cleaned_lines.append("")
            prev_blank = True
            continue
        line = re.sub(r"[ \t]+", " ", line)
        cleaned_lines.append(line)
        prev_blank = False

    return "\n".join(cleaned_lines).strip()


def is_placeholder_notes(text: str) -> bool:
    normalized = text.strip()
    if not normalized:
        return True
    return any(pattern.match(normalized) for pattern in PLACEHOLDER_PATTERNS)


def extract_slide_content(slide) -> str:
    """Collect visible text from slide shapes."""
    parts: List[str] = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if text:
            parts.append(text)
    return "\n\n".join(parts).strip()


def write_speaker_notes(slide, text: str) -> None:
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = text


def generate_narration_text(
    client: OpenAI,
    deck_name: str,
    slide_num: int,
    content: str,
    existing_notes: Optional[str] = None,
) -> str:
    if not content.strip():
        raise ValueError("no extractable slide content")

    existing_section = ""
    if existing_notes and not is_placeholder_notes(existing_notes):
        existing_section = f"\nExisting notes (optional hint):\n{existing_notes}\n"

    user_prompt = NARRATION_USER_TEMPLATE.format(
        deck_name=deck_name,
        slide_num=slide_num,
        content=content,
        existing_notes_section=existing_section,
    )
    response = client.chat.completions.create(
        model=NARRATION_MODEL,
        messages=[
            {"role": "system", "content": NARRATION_SYSTEM_PROMPT},
            {"role": "user", "content": user_prompt},
        ],
        temperature=0.7,
    )
    narration = (response.choices[0].message.content or "").strip()
    if not narration:
        raise ValueError("empty narration returned from model")
    if len(narration) > MAX_NARRATION_CHARS:
        narration = narration[:MAX_NARRATION_CHARS].rsplit(" ", 1)[0] + "."
    return narration


def parse_slide_spec(spec: str, total_slides: int) -> Set[int]:
    """Parse '5', '1,3,5', '1-10', or 'all' into slide numbers."""
    spec = spec.strip().lower()
    if not spec or spec == "all":
        return set(range(1, total_slides + 1))

    slides: Set[int] = set()
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            start_s, end_s = part.split("-", 1)
            start, end = int(start_s), int(end_s)
            if start > end:
                start, end = end, start
            slides.update(range(start, end + 1))
        else:
            slides.add(int(part))

    invalid = [n for n in slides if n < 1 or n > total_slides]
    if invalid:
        raise ValueError(
            f"Slide numbers out of range (1-{total_slides}): {sorted(invalid)}"
        )
    return slides


def audio_dir_for_pptx(pptx_path: Path) -> Path:
    return pptx_path.parent / f"{pptx_path.stem}_audio"


def split_tts_chunks(text: str, limit: int = MAX_TTS_INPUT_CHARS) -> List[str]:
    """Split long narration into TTS-safe chunks at sentence boundaries."""
    text = text.strip()
    if len(text) <= limit:
        return [text]

    chunks: List[str] = []
    current = ""
    parts = re.split(r"(?<=[.!?])\s+", text)
    for part in parts:
        if not part:
            continue
        candidate = f"{current} {part}".strip() if current else part
        if len(candidate) <= limit:
            current = candidate
            continue
        if current:
            chunks.append(current)
        if len(part) <= limit:
            current = part
        else:
            words = part.split()
            current = ""
            for word in words:
                candidate = f"{current} {word}".strip() if current else word
                if len(candidate) <= limit:
                    current = candidate
                else:
                    if current:
                        chunks.append(current)
                    current = word
    if current:
        chunks.append(current)
    return chunks


def concat_mp3_files(part_paths: List[Path], output_path: Path) -> None:
    """Concatenate MP3 parts with ffmpeg."""
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False) as tmp:
        for part in part_paths:
            tmp.write(f"file '{part}'\n")
        list_path = tmp.name
    try:
        result = subprocess.run(
            [
                "ffmpeg",
                "-y",
                "-f",
                "concat",
                "-safe",
                "0",
                "-i",
                list_path,
                "-c",
                "copy",
                str(output_path),
            ],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            raise RuntimeError(result.stderr or result.stdout or "ffmpeg concat failed")
    finally:
        Path(list_path).unlink(missing_ok=True)


def audio_path_for_slide(audio_dir: Path, slide_number: int) -> Path:
    return audio_dir / f"slide_{slide_number:03d}.{AUDIO_FORMAT}"


@dataclass
class SlideResult:
    slide_number: int
    narration: Optional[str] = None
    audio_path: Optional[Path] = None
    embedded: bool = False
    timed: bool = False
    generated: bool = False
    message: str = ""
    error: bool = False


@dataclass
class DeckSession:
    pptx_path: Path
    skip_backup: bool = False
    audio_dir: Path = field(init=False)
    _client: Optional[OpenAI] = field(default=None, init=False, repr=False)
    _backed_up: bool = field(default=False, init=False, repr=False)

    def __post_init__(self) -> None:
        self.pptx_path = self.pptx_path.expanduser().resolve()
        if not self.pptx_path.exists():
            raise FileNotFoundError(f"PowerPoint not found: {self.pptx_path}")
        if self.pptx_path.suffix.lower() != ".pptx":
            raise ValueError(f"Not a .pptx file: {self.pptx_path}")
        self.audio_dir = audio_dir_for_pptx(self.pptx_path)
        self.audio_dir.mkdir(parents=True, exist_ok=True)

    @property
    def client(self) -> OpenAI:
        if self._client is None:
            load_dotenv(BASE_DIR / ".env")
            api_key = os.environ.get("OPENAI_API_KEY")
            if not api_key or api_key == "your-api-key-here":
                raise RuntimeError("OPENAI_API_KEY is not set in .env")
            self._client = OpenAI()
        return self._client

    def load_presentation(self) -> Presentation:
        return Presentation(str(self.pptx_path))

    def slide_count(self) -> int:
        return len(self.load_presentation().slides)

    def get_narration(self, slide_number: int) -> Optional[str]:
        presentation = self.load_presentation()
        slide = presentation.slides[slide_number - 1]
        raw = ""
        if slide.has_notes_slide:
            text_frame = slide.notes_slide.notes_text_frame
            if text_frame is not None:
                raw = text_frame.text or ""
        cleaned = clean_narration(raw)
        if not cleaned or is_placeholder_notes(cleaned):
            return None
        return cleaned

    def generate_audio(
        self,
        slide_number: int,
        narration: Optional[str] = None,
        regenerate: bool = True,
    ) -> SlideResult:
        result = SlideResult(slide_number=slide_number)
        text = narration if narration is not None else self.get_narration(slide_number)
        if not text:
            result.message = "no narration"
            return result

        result.narration = text
        output_path = audio_path_for_slide(self.audio_dir, slide_number)
        if output_path.exists() and not regenerate:
            result.audio_path = output_path
            result.message = "audio already exists"
            return result

        chunks = split_tts_chunks(text)
        try:
            if len(chunks) == 1:
                with self.client.audio.speech.with_streaming_response.create(
                    model=TTS_MODEL,
                    voice=VOICE,
                    input=chunks[0],
                    instructions=SPEECH_INSTRUCTIONS,
                    response_format=AUDIO_FORMAT,
                ) as response:
                    response.stream_to_file(output_path)
                result.message = f"audio generated -> {output_path.name}"
            else:
                part_paths: List[Path] = []
                with tempfile.TemporaryDirectory(prefix="tts_parts_") as tmp:
                    tmp_dir = Path(tmp)
                    for idx, chunk in enumerate(chunks, start=1):
                        part_path = tmp_dir / f"part_{idx:02d}.{AUDIO_FORMAT}"
                        with self.client.audio.speech.with_streaming_response.create(
                            model=TTS_MODEL,
                            voice=VOICE,
                            input=chunk,
                            instructions=SPEECH_INSTRUCTIONS,
                            response_format=AUDIO_FORMAT,
                        ) as response:
                            response.stream_to_file(part_path)
                        part_paths.append(part_path)
                    concat_mp3_files(part_paths, output_path)
                result.message = (
                    f"audio generated ({len(chunks)} parts) -> {output_path.name}"
                )
            result.audio_path = output_path
        except Exception as exc:
            result.error = True
            result.message = str(exc)
        return result

    def _backup_if_needed(self) -> None:
        if self.skip_backup or self._backed_up:
            return
        backup_path = self.pptx_path.with_name(f"{self.pptx_path.stem}_bak.pptx")
        if not backup_path.exists():
            shutil.copy2(self.pptx_path, backup_path)
            print(f"Backup created: {backup_path.name}")
        self._backed_up = True

    def _set_autoplay(self, media_shape) -> bool:
        element = media_shape._element
        shape_id = xpath(element, ".//p:cNvPr")[0].attrib["id"]
        slide_element = element.getparent().getparent().getparent()
        for media_tag in ("video", "audio"):
            targets = xpath(
                slide_element,
                f'.//p:timing//p:{media_tag}//p:spTgt[@spid="{shape_id}"]',
            )
            if not targets:
                continue
            cond_nodes = xpath(targets[0].getparent().getparent(), ".//p:cond")
            if cond_nodes:
                cond_nodes[0].set("delay", "0")
                return True
        return False

    def _hide_media_icon(self, media_shape) -> bool:
        element = media_shape._element
        shape_id = xpath(element, ".//p:cNvPr")[0].attrib["id"]
        slide_element = element.getparent().getparent().getparent()
        for media_tag in ("video", "audio"):
            nodes = xpath(
                slide_element,
                f'.//p:timing//p:{media_tag}//p:spTgt[@spid="{shape_id}"]/ancestor::p:cMediaNode',
            )
            if nodes:
                nodes[0].set("showWhenStopped", "0")
                return True
        return False

    def embed_audio(self, slide_number: int, presentation: Presentation) -> SlideResult:
        result = SlideResult(slide_number=slide_number)
        audio_path = audio_path_for_slide(self.audio_dir, slide_number)
        if not audio_path.exists():
            result.message = "no MP3 — generate audio first"
            return result

        slide = presentation.slides[slide_number - 1]
        removed = remove_slide_media(slide)
        if removed:
            print(f"  removed {removed} existing media object(s) on slide {slide_number}")

        left = presentation.slide_width - ICON_MARGIN - ICON_SIZE
        top = presentation.slide_height - ICON_MARGIN - ICON_SIZE

        try:
            _patch_media_sha1_lookup()
            media = slide.shapes.add_movie(
                str(audio_path),
                left,
                top,
                ICON_SIZE,
                ICON_SIZE,
                mime_type="audio/mpeg",
            )
            if AUTO_PLAY:
                self._set_autoplay(media)
            if HIDE_ICON_DURING_SHOW:
                self._hide_media_icon(media)
            result.embedded = True
            result.message = f"embedded {audio_path.name}"
        except Exception as exc:
            result.error = True
            result.message = str(exc)
        return result

    def _advance_duration_ms(self, slide_number: int) -> int:
        audio_path = audio_path_for_slide(self.audio_dir, slide_number)
        if not audio_path.exists():
            return DEFAULT_SLIDE_DURATION_MS
        return int(MP3(str(audio_path)).info.length * 1000) + BUFFER_MS

    def set_timing(self, slide_number: int, slide) -> SlideResult:
        result = SlideResult(slide_number=slide_number)
        try:
            duration_ms = self._advance_duration_ms(slide_number)
            slide_element = slide._element
            adv_click = "1" if ALLOW_CLICK_ADVANCE else "0"
            transitions = xpath(slide_element, ".//p:transition")
            if transitions:
                transition = transitions[0]
            else:
                clr_map_nodes = xpath(slide_element, ".//p:clrMapOvr")
                if not clr_map_nodes:
                    raise ValueError("could not find p:clrMapOvr on slide")
                parent = clr_map_nodes[0].getparent()
                transition = etree.Element(
                    etree.QName(P_NS, "transition"),
                    attrib={"spd": "med"},
                )
                parent.insert(-1, transition)
            transition.set("advTm", str(duration_ms))
            transition.set("advClick", adv_click)
            transition.set("spd", "med")
            result.timed = True
            result.message = f"auto-advance after {duration_ms / 1000:.1f}s"
        except Exception as exc:
            result.error = True
            result.message = str(exc)
        return result

    def process_slides(
        self,
        slide_numbers: Iterable[int],
        *,
        show_narration: bool = False,
        generate_narration: bool = False,
        generate_audio: bool = False,
        embed: bool = False,
        set_timing: bool = False,
        regenerate_audio: bool = True,
    ) -> List[SlideResult]:
        numbers = sorted(set(slide_numbers))
        results: List[SlideResult] = []
        needs_save = generate_narration or embed or set_timing

        if needs_save:
            self._backup_if_needed()

        presentation = self.load_presentation() if needs_save else None
        deck_name = self.pptx_path.stem

        for slide_number in numbers:
            narration = self.get_narration(slide_number)

            if generate_narration and presentation is not None:
                slide = presentation.slides[slide_number - 1]
                content = extract_slide_content(slide)
                existing = narration
                try:
                    if not content.strip():
                        result = SlideResult(
                            slide_number=slide_number,
                            message="no slide content — skipped",
                        )
                        results.append(result)
                        print(f"Slide {slide_number}: {result.message}")
                        continue
                    narration = generate_narration_text(
                        self.client,
                        deck_name,
                        slide_number,
                        content,
                        existing_notes=existing,
                    )
                    write_speaker_notes(slide, narration)
                    gen_result = SlideResult(
                        slide_number=slide_number,
                        narration=narration,
                        generated=True,
                        message="narration generated",
                    )
                    results.append(gen_result)
                    print(f"Slide {slide_number}: {gen_result.message}")
                    if show_narration:
                        print(f"\n--- Slide {slide_number} ---")
                        print(narration)
                except Exception as exc:
                    result = SlideResult(
                        slide_number=slide_number,
                        error=True,
                        message=str(exc),
                    )
                    results.append(result)
                    print(f"Slide {slide_number}: ERROR — {exc}")
                    continue

            if show_narration and not generate_narration:
                result = SlideResult(slide_number=slide_number, narration=narration)
                if narration:
                    result.message = "narration found"
                else:
                    result.message = "no narration"
                results.append(result)
                print(f"\n--- Slide {slide_number} ---")
                if narration:
                    print(narration)
                else:
                    print("(no narration)")

            if generate_audio:
                audio_result = self.generate_audio(
                    slide_number,
                    narration=narration,
                    regenerate=regenerate_audio,
                )
                if not show_narration and not generate_narration:
                    results.append(audio_result)
                elif generate_narration:
                    audio_result.narration = narration
                    results[-1] = audio_result if results else audio_result
                print(f"Slide {slide_number}: {audio_result.message}")

            if embed and presentation is not None:
                embed_result = self.embed_audio(slide_number, presentation)
                if not show_narration and not generate_audio and not generate_narration:
                    results.append(embed_result)
                print(f"Slide {slide_number}: {embed_result.message}")

            if set_timing and presentation is not None:
                slide = presentation.slides[slide_number - 1]
                timing_result = self.set_timing(slide_number, slide)
                if (
                    not show_narration
                    and not generate_audio
                    and not embed
                    and not generate_narration
                ):
                    results.append(timing_result)
                print(f"Slide {slide_number}: {timing_result.message}")

        if needs_save and presentation is not None:
            safe_save_presentation(presentation, self.pptx_path)
            print(f"\nSaved: {self.pptx_path}")

        return results
